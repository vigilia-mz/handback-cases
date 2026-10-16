"""End-to-end dry run with mock models. No API key, no cost.

    python scripts/dry_run.py

Runs every step of the pipeline against mock models and writes all output to
logs/dryrun/ (never to ratings/):
  1. generation with --no-score (mock model under test)
  2. the blinded workbook, then fake "author" scores filled in by script
  3. scoring the generation log with a mock grader that returns valid JSON
  4. scoring the constructed replies, plus a deliberately broken grader to check
     that unparseable output is recorded as unscored, never guessed
  5. the gaming report and the agreement report
The numbers produced are meaningless by design: this checks plumbing only.
"""
from __future__ import annotations

import json
import random
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "logs" / "dryrun"
sys.path.insert(0, str(ROOT))

# Offline token counting: mockllm counts tokens with tiktoken, which downloads an
# encoding file on first use. Replace it with a whitespace counter for this dry run only.
import tiktoken  # noqa: E402


class _Enc:
    def encode(self, s, **k):
        return s.split()


tiktoken.get_encoding = lambda name: _Enc()
tiktoken.encoding_for_model = lambda name: _Enc()

from inspect_ai import eval as inspect_eval, score as inspect_score  # noqa: E402
from inspect_ai.model import ModelOutput, get_model  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from evals.handback import handback, handback_constructed  # noqa: E402
from evals.scorers import criteria_scorer, dimension_scorer, label_scorer  # noqa: E402


def mock_reply(messages, tools, tool_choice, config):
    last = messages[-1].text[:40]
    return ModelOutput.from_content(model="mockllm", content=f"[mock reply to: {last}...] I can help with that.")


def mock_grader(rng):
    def grade(messages, tools, tool_choice, config):
        p = messages[-1].text
        if "score the reply on three dimensions" in p:
            body = {k: {"quote": "mock quote", "level": rng.randint(0, 3), "reason": "mock"}
                    for k in ["contextual_judgment", "usefulness", "paternalism"]}
        elif "Apply every label" in p:
            body = {"labels": [{"id": i, "quote": "mock"} for i in sorted(rng.sample(range(1, 6), 2))]}
        else:
            body = {"met": rng.random() < 0.5, "quote": "mock", "reason": "mock"}
        return ModelOutput.from_content(model="mockllm", content="```json\n" + json.dumps(body) + "\n```")
    return grade


def broken_grader(messages, tools, tool_choice, config):
    return ModelOutput.from_content(model="mockllm", content="I think it is a 2, roughly.")


def run(cmd):
    print("$", " ".join(cmd))
    r = subprocess.run(cmd, cwd=ROOT, capture_output=True, text=True)
    print(r.stdout[-1500:])
    if r.returncode:
        print(r.stderr[-3000:])
        raise SystemExit(f"step failed: {cmd[1]}")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    rng = random.Random(7)

    run([sys.executable, "scripts/validate_cases.py", "--allow-drafts"])

    print("\n1. Generation, no scoring")
    under_test = get_model("mockllm/model", custom_outputs=mock_reply)
    [gen] = inspect_eval(handback(allow_drafts=True), model=under_test, epochs=2, score=False,
                         log_dir=str(OUT), display="none")
    assert gen.status == "success", gen.status
    assert all(not s.scores for s in gen.samples), "generation log must hold no scores"
    print(f"   {len(gen.samples)} replies, no scores in the log")

    print("\n2. Blinded workbook")
    run([sys.executable, "scripts/export_for_raters.py", gen.location, "--rater", "rater1", "--out-dir", str(OUT)])
    wb = load_workbook(OUT / "rater1_scoring.xlsx")
    ws = wb["Score"]
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "mockllm" not in text.replace("[mock reply", ""), "model name leaked into the workbook"
    for row in ws.iter_rows(min_row=2):
        for col in range(6, 9):
            row[col - 1].value = rng.randint(0, 3)
        for col in range(9, 14):
            row[col - 1].value = rng.choice(["Y", "N"])
        for col in range(14, 21):
            if row[col - 1].value != "n/a":
                row[col - 1].value = rng.choice(["Y", "N"])
        row[20].value = "mock reason"
    wb.save(OUT / "rater1_scoring.xlsx")
    print("   workbook has", ws.max_row - 1, "rows; fake author scores filled in")

    print("\n3. Scoring the generation log with a mock grader")
    grader = get_model("mockllm/model", custom_outputs=mock_grader(rng))
    scored = inspect_score(gen, scorers=[dimension_scorer(), label_scorer(), criteria_scorer()],
                           model_roles={"grader": grader}, display="none")
    scored_path = OUT / "scored_generation.eval"
    from inspect_ai.log import write_eval_log
    write_eval_log(scored, str(scored_path))
    s0 = scored.samples[0].scores
    print("   scorers:", sorted(s0))
    print("   first sample:", {k: v.value for k, v in s0.items()})

    print("\n4. Constructed replies, mock grader, then a broken grader")
    [con] = inspect_eval(handback_constructed(allow_drafts=True), model=get_model("mockllm/model"),
                         model_roles={"grader": get_model("mockllm/model", custom_outputs=mock_grader(rng))},
                         log_dir=str(OUT), display="none")
    assert con.status == "success", con.status
    assert all(s.output.completion.startswith(("From", "Because", "Before")) for s in con.samples), \
        "constructed replies were not used verbatim"
    print(f"   {len(con.samples)} constructed replies scored; text used verbatim")
    broken = inspect_score(con, scorers=[dimension_scorer(), label_scorer(), criteria_scorer()],
                           model_roles={"grader": get_model("mockllm/model", custom_outputs=broken_grader)},
                           action="overwrite", display="none")
    b0 = broken.samples[0].scores
    assert b0["dimension_scorer"].value == {} and b0["criteria_scorer"].value == {}, b0
    print("   broken grader output recorded as unscored:", b0["dimension_scorer"].explanation)
    broken_path = OUT / "scored_constructed_brokengrader.eval"
    write_eval_log(broken, str(broken_path))

    print("\n5. Reports")
    run([sys.executable, "scripts/gaming_report.py", con.location, str(broken_path), "--out", str(OUT / "gaming_report.md")])
    run([sys.executable, "scripts/agreement.py", "--rater", str(OUT / "rater1_scoring.xlsx"),
         "--key", str(OUT / "rater1_key.json"), "--grader", str(scored_path), "--out", str(OUT / "agreement_report.md")])

    print("\n6. Real-grader guard: scoring without a bound grader must fail")
    try:
        inspect_score(gen, scorers=[dimension_scorer()], display="none")
        raise SystemExit("FAILED: scoring ran without a grader role")
    except SystemExit:
        raise
    except Exception as e:  # noqa: BLE001
        print("   refused as expected:", type(e).__name__)

    print("\nDry run complete. Everything is in logs/dryrun/. Nothing was written to ratings/.")


if __name__ == "__main__":
    main()
