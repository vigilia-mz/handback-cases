"""Write the blinded scoring workbook from a generation log.

    python scripts/export_for_raters.py logs/<generation>.eval --rater rater1 [--seed 20261016]

Writes ratings/<rater>_scoring.xlsx (three tabs: Score, Cases, Anchors) and
ratings/<rater>_key.json, which maps row ids back to samples. The workbook holds no
model name and no grader scores. Rows are shuffled with a recorded seed.
"""
from __future__ import annotations

import argparse
import json
import random
import sys
from pathlib import Path

import yaml
from inspect_ai.log import read_eval_log
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from evals.dataset import load_rubric  # noqa: E402

MAX_CRIT = 7
WRAP = Alignment(wrap_text=True, vertical="top")
GREY = PatternFill("solid", fgColor="DDDDDD")
HEAD = Font(bold=True)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("log")
    ap.add_argument("--rater", default="rater1")
    ap.add_argument("--seed", type=int, default=20261016)
    ap.add_argument("--out-dir", default=str(ROOT / "ratings"))
    a = ap.parse_args()

    log = read_eval_log(a.log)
    if any(s.scores for s in (log.samples or [])):
        sys.exit("This log already has scores. Export from the generation log made with --no-score.")
    cases = {p.stem: yaml.safe_load(p.read_text(encoding="utf-8")) for p in (ROOT / "cases/gold").glob("*.yaml")}
    rubric = load_rubric()

    items = [(s.id, s.epoch, s.output.completion if s.output else "") for s in log.samples]
    rng = random.Random(a.seed)
    rng.shuffle(items)

    wb = Workbook()
    ws = wb.active
    ws.title = "Score"
    lab = {l["id"]: l["name"] for l in rubric["labels"]}
    header = (["Row", "Case", "Final user turn", "Reply", "Criteria for this case (C1, C2, ...)",
               "Contextual judgment (0-3)", "Usefulness (0-3)", "Paternalism (0-3, 3 = none)"]
              + [f"L{i} {lab[i]} (Y/N)" for i in range(1, 6)]
              + [f"C{i} met (Y/N)" for i in range(1, MAX_CRIT + 1)]
              + ["Reason (quote the reply)", "Notes (write ANCHOR? if an anchor did not fit)"])
    ws.append(header)
    for cell in ws[1]:
        cell.font, cell.alignment = HEAD, WRAP

    dv_level = DataValidation(type="list", formula1='"0,1,2,3"', allow_blank=True)
    dv_yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    ws.add_data_validation(dv_level)
    ws.add_data_validation(dv_yn)

    key = []
    for n, (sid, epoch, reply) in enumerate(items, 1):
        row_id = f"X{n:02d}"
        c = cases[sid]
        crit = "\n".join(f"C{i}. ({k['points']:+d}) {k['text']}" for i, k in enumerate(c["criteria"], 1))
        ws.append([row_id, sid, c["final_user"], reply, crit] + [None] * (3 + 5 + MAX_CRIT + 2))
        r = ws.max_row
        for col in range(6, 9):
            dv_level.add(ws.cell(r, col))
        for col in range(9, 14):
            dv_yn.add(ws.cell(r, col))
        for i in range(1, MAX_CRIT + 1):
            cell = ws.cell(r, 13 + i)
            if i <= len(c["criteria"]):
                dv_yn.add(cell)
            else:
                cell.value, cell.fill = "n/a", GREY
        for cell in ws[r]:
            cell.alignment = WRAP
        ws.row_dimensions[r].height = 400
        key.append({"row": row_id, "sample_id": sid, "epoch": epoch})

    widths = {"A": 6, "B": 6, "C": 30, "D": 80, "E": 45}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for i in range(6, len(header) + 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = 14
    ws.column_dimensions[ws.cell(1, len(header) - 1).column_letter].width = 45
    ws.column_dimensions[ws.cell(1, len(header)).column_letter].width = 30
    ws.freeze_panes = "C2"

    wc = wb.create_sheet("Cases")
    wc.append(["Case", "Type", "Conversation (scripted turns, then the final user turn)", "What the judgment rests on", "Values in their own words"])
    for cell in wc[1]:
        cell.font = HEAD
    for cid in sorted(cases):
        c = cases[cid]
        conv = "\n\n".join(f"USER: {t['user']}\n\nASSISTANT: {t['assistant']}" for t in c["turns"])
        conv += f"\n\nUSER (final turn): {c['final_user']}"
        vals = "\n".join(f"(turn {v['turn']}) {v['text']}" for v in c["values_stated"])
        wc.append([cid, c["type"], conv, c.get("judgment_rests_on"), vals])
        for cell in wc[wc.max_row]:
            cell.alignment = WRAP
    for col, w in {"A": 6, "B": 9, "C": 100, "D": 45, "E": 45}.items():
        wc.column_dimensions[col].width = w

    wa = wb.create_sheet("Anchors")
    wa.append(["Dimension", "Level", "Anchor"])
    for cell in wa[1]:
        cell.font = HEAD
    for k, d in rubric["dimensions"].items():
        for lvl in [3, 2, 1, 0]:
            wa.append([d["title"], lvl, d["levels"][lvl]])
            wa.cell(wa.max_row, 3).alignment = WRAP
    wa.append([])
    wa.append(["Labels (a reply can carry several)"])
    wa.cell(wa.max_row, 1).font = HEAD
    for l in rubric["labels"]:
        wa.append([f"L{l['id']} {l['name']}", None, l["definition"]])
        wa.cell(wa.max_row, 3).alignment = WRAP
    for col, w in {"A": 28, "B": 6, "C": 100}.items():
        wa.column_dimensions[col].width = w

    out_dir = Path(a.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"{a.rater}_scoring.xlsx"
    wb.save(out)
    (out_dir / f"{a.rater}_key.json").write_text(json.dumps(
        {"log": Path(a.log).name, "seed": a.seed, "rows": key}, indent=1), encoding="utf-8")
    print(f"wrote {out} ({len(items)} rows, seed {a.seed}) and the row key")


if __name__ == "__main__":
    main()
